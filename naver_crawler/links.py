"""링크 목록 파일(txt / csv / xlsx) 읽기."""

from __future__ import annotations

from pathlib import Path

from naver_crawler.parsing import extract_urls

SUPPORTED = (".txt", ".csv", ".tsv", ".xlsx", ".xlsm")


def read_text_file(path: str | Path) -> str:
    """UTF-8, UTF-16(BOM), 메모장 ANSI(cp949) 를 자동으로 구분해 읽는다."""
    data = Path(path).read_bytes()
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _xlsx_text(path: str | Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.hyperlink is not None and cell.hyperlink.target:
                    lines.append(str(cell.hyperlink.target))
                elif cell.value is not None:
                    lines.append(str(cell.value))
    return "\n".join(lines)


def load_links_file(path: str | Path) -> list[str]:
    """파일 안의 링크만 순서대로 돌려준다 (중복 제거 안 함)."""
    suffix = Path(path).suffix.lower()
    text = _xlsx_text(path) if suffix in (".xlsx", ".xlsm") else read_text_file(path)
    return extract_urls(text)


def dedupe(urls: list[str], existing: list[str] | None = None) -> tuple[list[str], int]:
    seen = set(existing or [])
    out = []
    for url in urls:
        if url not in seen:
            seen.add(url)
            out.append(url)
    return out, len(urls) - len(out)
