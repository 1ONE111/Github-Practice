import datetime as dt

from openpyxl import load_workbook

from naver_crawler.crawlers import TASK_BY_ID
from naver_crawler.excel import FIRST_DATA_ROW, HEADER_ROW, TITLE_ROW, SheetData, default_filename, export_xlsx

WHEN = dt.datetime(2026, 10, 7, 14, 30)


def _cafe_sheet():
    task = TASK_BY_ID["cafe_article"]
    rows = [
        {"url": "https://cafe.naver.com/a/1", "title": "첫 글", "nickname": "요리왕", "views": 1234, "comments": 5, "likes": 3,
         "date": "2026.10.01. 12:00", "status": "완료", "message": ""},
        {"url": "https://cafe.naver.com/a/2", "title": None, "views": None, "comments": None, "likes": None,
         "status": "실패", "message": "값을 찾지 못했습니다"},
        {"url": "https://cafe.naver.com/a/3", "title": "셋째", "views": 10, "comments": 1, "likes": None,
         "status": "일부 누락", "message": "누락: 좋아요"},
    ]
    return SheetData(task.title, task.description, task.columns, rows, task.notes)


def test_list_layout_matches_team_template(tmp_path):
    wb = load_workbook(export_xlsx(tmp_path / "out", [_cafe_sheet()], WHEN))
    assert wb.sheetnames == ["카페 게시글"]
    ws = wb["카페 게시글"]

    # B2 병합 제목: 검정 바탕 흰 글씨 '제목 (건수)'
    assert ws.cell(TITLE_ROW, 2).value == "카페 게시글 (3)"
    assert f"B{TITLE_ROW}:J{TITLE_ROW}" in {str(r) for r in ws.merged_cells.ranges}
    assert ws.cell(TITLE_ROW, 2).fill.fgColor.rgb.endswith("000000")
    assert ws.cell(TITLE_ROW, 2).font.color.rgb.endswith("FFFFFF")
    assert ws.column_dimensions["A"].width == 2.8

    headers = [ws.cell(HEADER_ROW, c).value for c in range(2, 11)]
    assert headers == ["구분", "제목", "URL", "닉네임", "조회수", "댓글수", "좋아요", "작성일", "비고"]
    assert ws.cell(HEADER_ROW, 2).fill.fgColor.rgb.endswith("000000")
    assert ws.row_dimensions[HEADER_ROW + 1].height == 3.75

    r = FIRST_DATA_ROW
    assert [ws.cell(r, c).value for c in range(2, 11)] == [
        1, "첫 글", "https://cafe.naver.com/a/1", "요리왕", 1234, 5, 3, "2026.10.01. 12:00", None]
    assert ws.cell(r, 6).number_format == "#,##0\\ "
    assert ws.cell(r, 4).hyperlink.target == "https://cafe.naver.com/a/1"
    assert ws.cell(r, 4).font.color.rgb.endswith("0000FF")
    assert ws.cell(r, 3).font.size == 8 and ws.cell(r, 3).border.left.style == "thin"
    assert ws.cell(r, 6).alignment.horizontal == "left"
    # 실패/일부 누락은 색 대신 비고에 글로
    assert ws.cell(r + 1, 3).value == "-" and ws.cell(r + 1, 6).value == "-"
    assert ws.cell(r + 1, 10).value == "실패 - 값을 찾지 못했습니다"
    assert ws.cell(r + 2, 10).value == "누락: 좋아요"
    # 합계/평균 행 없이 한 줄 띄고 주석
    assert ws.cell(r + 3, 2).value is None
    assert ws.cell(r + 4, 2).value == "* 수집일시: 2026-10-07 14:30"


def test_multi_sheet_no_summary(tmp_path):
    other = TASK_BY_ID["kin"]
    sheets = [_cafe_sheet(), SheetData(other.title, other.description, other.columns, [])]
    wb = load_workbook(export_xlsx(tmp_path / "all.xlsx", sheets, WHEN))
    assert wb.sheetnames == ["카페 게시글", "지식iN"]
    assert wb["지식iN"].cell(TITLE_ROW, 2).value == "지식iN (0)"


def test_blog_sheet_order_sort_and_no_widget_exclusion(tmp_path):
    task = TASK_BY_ID["blog_info"]
    rows = [
        {"url": "https://blog.naver.com/a", "name": "A", "nickname": "a", "avg": 100, "buddies": 5868, "status": "완료"},
        {"url": "https://blog.naver.com/b", "name": "B", "nickname": "b", "avg": None, "buddies": 12,
         "status": "위젯없음", "message": "방문자 위젯이 없는 블로그"},
        {"url": "https://blog.naver.com/c", "name": "C", "nickname": "c", "avg": 300, "buddies": 319000, "status": "완료"},
        {"url": "https://blog.naver.com/d", "name": "D", "nickname": "d", "avg": 50, "buddies": None,
         "status": "일부 누락", "message": "누락: 이웃수"},
    ]
    sheet = SheetData(task.title, task.description, task.columns, rows, task.notes, task.sort_desc)
    wb = load_workbook(export_xlsx(tmp_path / "blog.xlsx", [sheet], WHEN))
    assert wb.sheetnames == ["블로그 정보", "제외 목록"]
    ws = wb["블로그 정보"]
    assert ws.cell(TITLE_ROW, 2).value == "블로그 정보 (3)"
    headers = [ws.cell(HEADER_ROW, c).value for c in range(2, 9)]
    assert headers == ["구분", "닉네임", "블로그명", "블로그링크", "이웃수", "일방문자수(5일 평균)", "비고"]
    # 이웃 많은 순, 이웃수 없는 행은 맨 뒤, 구분은 정렬 후 1부터
    body = [[ws.cell(r, c).value for c in range(2, 9)] for r in range(FIRST_DATA_ROW, FIRST_DATA_ROW + 3)]
    assert body == [
        [1, "c", "C", "https://blog.naver.com/c", 319000, 300, None],
        [2, "a", "A", "https://blog.naver.com/a", 5868, 100, None],
        [3, "d", "D", "https://blog.naver.com/d", "-", 50, "누락: 이웃수"],
    ]
    assert ws.cell(FIRST_DATA_ROW, 5).hyperlink.target == "https://blog.naver.com/c"
    notes = [ws.cell(r, 2).value for r in range(FIRST_DATA_ROW + 4, FIRST_DATA_ROW + 8)]
    assert notes[1].startswith("* 이웃수:")
    assert notes[2].startswith("* 일방문자수(5일 평균):")
    assert "1건은 '제외 목록'" in notes[3]

    excluded = wb["제외 목록"]
    assert excluded.cell(TITLE_ROW, 2).value == "제외 목록 (1)"
    assert [excluded.cell(FIRST_DATA_ROW, c).value for c in range(2, 7)] == [
        1, "블로그 정보", "B", "https://blog.naver.com/b", "방문자 위젯이 없는 블로그"]


def test_default_filename():
    assert default_filename("카페 게시글", dt.datetime(2026, 10, 6, 3, 4)) == "네이버크롤링_카페_게시글_261006.xlsx"
