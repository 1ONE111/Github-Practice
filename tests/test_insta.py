"""인스타 크롤러: 링크 해석, 응답 → 평균 계산, 요청 제한 · 로그인 처리, 엑셀 저장 (네트워크 없이)."""

import json
import threading

import pytest
from openpyxl import load_workbook

from insta_crawler.client import InstaClient, InstaError, LoginRequired, NotFound, RateLimited
from insta_crawler.collector import COLUMNS, TITLES, Runner, StopRun, collect, notes
from insta_crawler.config import Settings, load_session, save_session
from insta_crawler.files import export_xlsx, read_links_text
from insta_crawler.parsing import parse_targets, recent_posts, posts_from_feed, shortcode_to_media_id, summarize

# ---------------------------------------------------------------- 링크 해석


def test_parse_targets_from_mixed_text():
    text = """
    # 메모 줄은 무시
    https://www.instagram.com/Some.User/?igsh=abc123
    인플루언서 instagram.com/second_user 팔로워 많음
    @third.one
    fourth
    https://instagram.com/stories/story_user/3141592653/
    https://www.instagram.com/_u/deep_link
    https://www.instagram.com/reel/DAbCdEfGhIj/?utm_source=ig_web
    https://www.instagram.com/explore/tags/food/
    문의 메일 contact@instagram.com
    한글만 있는 줄
    """
    targets = parse_targets(text)
    assert [t.username for t in targets] == [
        "some.user", "second_user", "third.one", "fourth", "story_user", "deep_link", None, None,
    ]
    assert targets[6].shortcode == "DAbCdEfGhIj"
    assert targets[7].shortcode is None  # explore 는 아이디가 아님


def test_bad_usernames_are_skipped():
    assert parse_targets(".dot\nend.\na..b\n" + "x" * 31) == []


def test_shortcode_to_media_id():
    # 64진수(A-Z a-z 0-9 - _) → 10진수. 2020년 4월 게시물 → ID 상위 비트가 그 시각
    media_id = shortcode_to_media_id("B-a8WxpH5Mw")
    assert media_id == "2277398009511646000"
    assert 2020.2 < 2011.65 + (int(media_id) >> 23) / 1000 / 86400 / 365.25 < 2020.4
    assert shortcode_to_media_id("B-a8WxpH5Mw" + "extra_private_suffix") == media_id
    assert shortcode_to_media_id("!!") is None


# ---------------------------------------------------------------- 가짜 인스타 응답


def node(code, ts, likes, comments, views=None, video=False, pinned=False):
    n = {"shortcode": code, "taken_at_timestamp": ts, "is_video": video,
         "edge_liked_by": {"count": likes}, "edge_media_to_comment": {"count": comments},
         "pinned_for_users": [{"id": "1"}] if pinned else []}
    if video:
        n["video_view_count"] = views
        n["product_type"] = "clips"
    return {"node": n}


def profile(username="alice", followers=12345, edges=None, private=False, total=None):
    edges = edges if edges is not None else []
    return {"data": {"user": {
        "id": "999", "username": username, "full_name": "앨리스", "is_private": private,
        "edge_followed_by": {"count": followers}, "edge_follow": {"count": 10},
        "edge_owner_to_timeline_media": {"count": len(edges) if total is None else total, "edges": edges},
    }}, "status": "ok"}


class Resp:
    def __init__(self, status=200, body="", headers=None):
        self.status_code = status
        self.text = body if isinstance(body, str) else json.dumps(body)
        self.headers = headers or {}


def fake_client(routes, cookies=None):
    """routes: {경로 앞부분: Resp 또는 [Resp, Resp…](차례대로)}"""
    client = InstaClient(cookies)
    calls = []

    def get(url, headers=None, timeout=None, allow_redirects=True, data=None):
        calls.append(url if data is None else ("POST", url, data))
        for prefix, resp in routes.items():
            if url.startswith("https://www.instagram.com" + prefix):
                if isinstance(resp, list):
                    return resp.pop(0) if len(resp) > 1 else resp[0]
                return resp
        return Resp(404, "")

    client.session.get = get
    client.session.post = get
    client.calls = calls
    return client


PROFILE = "/api/v1/users/web_profile_info/?username="
LOGGED_IN = [{"name": "sessionid", "value": "s"}, {"name": "csrftoken", "value": "c"}]


EDGES = [
    node("pinned_old", 1000, 99999, 9999, views=1, video=True, pinned=True),  # 고정된 옛 게시물: 날짜순으로 밀려나야 함
    node("a", 7000, 100, 10),
    node("b", 6000, 200, 20, views=5000, video=True),
    node("c", 5000, 300, 30),
    node("d", 4000, 400, 40, views=7000, video=True),
    node("e", 3000, 500, 50),
    node("f", 2000, 600, 60),
]


def test_feed_mode_latest_five_by_date_likes_and_comments_only():
    client = fake_client({PROFILE + "alice": Resp(200, profile(edges=EDGES))})
    res = collect(client, "alice", 5, "feed")
    assert res.status == "완료", res.message
    assert res.values == {
        "username": "alice", "name": "앨리스", "url": "https://www.instagram.com/alice/", "followers": 12345,
        "avg_likes": 300, "avg_comments": 30,
    }
    assert len(client.calls) == 1  # 비로그인은 프로필 요청 하나로 끝


def test_reels_mode_logged_out_uses_videos_among_recent_posts():
    client = fake_client({PROFILE + "alice": Resp(200, profile(edges=EDGES))})
    res = collect(client, "alice", 5, "reels")
    assert res.values == {
        "username": "alice", "name": "앨리스", "url": "https://www.instagram.com/alice/reels/", "followers": 12345,
        "avg_likes": 33533, "avg_comments": 3353, "avg_views": 4000,  # b · d · 고정된 옛 릴스 = 릴스 3개 전부
    }
    assert res.status == "완료" and "릴스 3개 기준 (비로그인" in res.message


def test_reels_mode_logged_in_uses_reels_tab():
    clips = {"items": [
        {"media": {"code": f"r{i}", "taken_at": 100 - i, "media_type": 2, "product_type": "clips",
                   "like_count": 10 * i, "comment_count": i, "play_count": 1000 * i}} for i in range(1, 8)
    ], "status": "ok"}
    client = fake_client({
        PROFILE + "bob": Resp(200, profile("bob", edges=EDGES)),
        "/api/v1/clips/user/": Resp(200, clips),
    }, LOGGED_IN)
    res = collect(client, "bob", 5, "reels")
    assert (res.values["avg_views"], res.values["avg_likes"], res.values["avg_comments"]) == (3000, 30, 3)
    assert res.status == "완료" and res.message == ""
    assert client.calls[1] == ("POST", "https://www.instagram.com/api/v1/clips/user/",
                               {"target_user_id": "999", "page_size": "12", "include_feed_video": "true"})


def test_feed_mode_logged_in_uses_feed_api():
    feed = {"items": [
        {"code": "r1", "taken_at": 300, "media_type": 2, "product_type": "clips", "like_count": 10,
         "comment_count": 1, "play_count": 1000, "view_count": 50},
        {"code": "p1", "taken_at": 200, "media_type": 1, "like_count": 20, "comment_count": 3},
        {"code": "c1", "taken_at": 100, "media_type": 8, "like_count": 30, "comment_count": 5},
    ], "status": "ok"}
    client = fake_client({
        PROFILE + "bob": Resp(200, profile("bob", edges=[node("x", 1, 1, 1)], total=3)),
        "/api/v1/feed/user/999/": Resp(200, feed),
    }, LOGGED_IN)
    res = collect(client, "bob", 5, "feed")
    assert res.values["avg_likes"] == 20 and res.values["avg_comments"] == 3 and "avg_views" not in res.values
    assert res.status == "완료" and res.message == "게시물 3개 기준"
    reels = collect(client, "bob", 5, "reels")  # 릴스 탭 API 실패 → 피드에서 영상만
    assert reels.values["avg_views"] == 1000 and reels.message == "릴스 1개 기준"


def test_private_empty_and_no_reels():
    client = fake_client({
        PROFILE + "secret": Resp(200, profile("secret", private=True, total=40)),
        PROFILE + "empty": Resp(200, profile("empty", total=0)),
        PROFILE + "photos": Resp(200, profile("photos", edges=[node("a", 1, 5, 1), node("b", 2, 7, 1)])),
    })
    private = collect(client, "secret", 5)
    assert private.status == "비공개" and private.values["followers"] == 12345
    empty = collect(client, "empty", 5)
    assert empty.status == "완료" and empty.message == "게시물 없음"
    assert collect(client, "photos", 5, "feed").values["avg_likes"] == 6
    photos = collect(client, "photos", 5, "reels")
    assert photos.values["avg_views"] is None and "릴스 없음" in photos.message


def test_hidden_like_counts_are_left_out_of_average():
    posts = posts_from_feed({"items": [
        {"taken_at": 2, "media_type": 1, "like_count": 100, "comment_count": 1},
        {"taken_at": 1, "media_type": 1, "comment_count": 1},
    ]})
    summary = summarize(recent_posts(posts, 5))
    assert summary.avg_likes == 100 and summary.hidden_likes == 1


def test_error_responses():
    client = fake_client({
        PROFILE + "gone": Resp(200, {"data": {"user": None}, "status": "ok"}),
        PROFILE + "slow": Resp(401, {"message": "Please wait a few minutes before you try again.",
                                     "require_login": True, "status": "fail"}),
        PROFILE + "login": Resp(302, "", {"Location": "https://www.instagram.com/accounts/login/"}),
        PROFILE + "html": Resp(200, "<html>로그인 페이지</html>"),
        PROFILE + "bio": Resp(200, profile("bio")),
    })
    with pytest.raises(NotFound):
        client.profile("gone")
    with pytest.raises(NotFound):
        client.profile("nobody")  # 404
    with pytest.raises(RateLimited):
        client.profile("slow")
    with pytest.raises(LoginRequired):
        client.profile("login")
    with pytest.raises(LoginRequired):
        client.profile("html")


def test_bio_words_do_not_look_like_errors():
    body = profile("bio")
    body["data"]["user"]["biography"] = "Please wait for my challenge! login_required spam"
    client = fake_client({PROFILE + "bio": Resp(200, body)})
    assert client.profile("bio")["username"] == "bio"


# ---------------------------------------------------------------- 실행기


def run(client, names, **settings):
    settings.setdefault("mode", "feed")
    runner = Runner(Settings(request_delay=0, cooldown=0, **settings), client)
    results = {}
    targets = parse_targets("\n".join(names))
    done = runner.run(list(enumerate(targets)), threading.Event(), on_result=results.__setitem__)
    return done, results


def test_runner_waits_once_on_rate_limit_then_continues():
    client = fake_client({
        PROFILE + "a": [Resp(429, ""), Resp(200, profile("a", edges=[node("x", 1, 10, 2)]))],
        PROFILE + "b": Resp(200, profile("b", edges=[node("y", 1, 20, 4)])),
    })
    done, results = run(client, ["a", "b"])
    assert done is True
    assert results[0].values["avg_likes"] == 10 and results[1].values["avg_likes"] == 20


def test_runner_stops_when_still_limited_or_login_needed():
    client = fake_client({PROFILE + "a": Resp(429, ""), PROFILE + "b": Resp(200, profile("b"))})
    with pytest.raises(StopRun, match="실패 재시도"):
        run(client, ["a", "b"])
    client = fake_client({PROFILE + "a": Resp(302, "", {"Location": "/accounts/login/"})})
    with pytest.raises(StopRun, match="로그인"):
        run(client, ["a"])


def test_runner_marks_single_failures_and_post_links():
    client = fake_client({PROFILE + "ok": Resp(200, profile("ok", edges=[node("x", 1, 3, 1)]))})
    done, results = run(client, ["nobody", "ok", "https://www.instagram.com/p/B-a8WxpH5Mw/"])
    assert done is True
    assert results[0].status == "실패" and "찾을 수 없습니다" in results[0].message
    assert results[1].status == "완료"
    assert results[2].status == "실패" and "로그인 후" in results[2].message


def test_post_link_resolves_owner_when_logged_in():
    client = fake_client({
        "/api/v1/media/2277398009511646000/info/": Resp(200, {"items": [{"user": {"username": "owner"}}]}),
        PROFILE + "owner": Resp(200, profile("owner", edges=[node("x", 1, 3, 1)])),
        "/api/v1/feed/user/999/": Resp(200, {"items": [], "status": "ok"}),
    }, LOGGED_IN)
    done, results = run(client, ["https://www.instagram.com/reel/B-a8WxpH5Mw/"])
    assert results[0].values["username"] == "owner" and results[0].values["avg_likes"] == 3


# ---------------------------------------------------------------- 파일


def test_session_roundtrip(tmp_path):
    path = tmp_path / "session.json"
    save_session([{"name": "sessionid", "value": "v", "domain": ".instagram.com", "path": "/", "httpOnly": True}], path)
    assert load_session(path) == [{"name": "sessionid", "value": "v", "domain": ".instagram.com", "path": "/"}]
    assert InstaClient(load_session(path)).logged_in


def test_read_links_and_export(tmp_path):
    src = tmp_path / "list.txt"
    src.write_bytes("@alpha\nhttps://instagram.com/beta\n".encode("cp949"))
    assert [t.username for t in parse_targets(read_links_text(src))] == ["alpha", "beta"]

    rows = [
        {"username": "alpha", "url": "https://www.instagram.com/alpha/reels/", "followers": 1500, "avg_likes": None,
         "avg_comments": None, "avg_views": None, "status": "완료", "message": "릴스 없음"},
        {"username": "beta", "url": "https://www.instagram.com/beta/", "status": "실패", "message": "계정을 찾을 수 없습니다"},
    ]
    out = export_xlsx(tmp_path / "out", TITLES["reels"], COLUMNS["reels"], rows, notes("reels", 5))
    ws = load_workbook(out).active
    assert ws["B2"].value == "인스타 릴스 지표 (2)"
    headers = [ws.cell(row=4, column=c).value for c in range(2, 2 + len(COLUMNS["reels"]) + 2)]
    assert headers == ["구분", "아이디", "이름", "릴스 링크", "팔로워", "평균 조회수", "평균 좋아요", "평균 댓글", "비고"]
    assert ws["F6"].value == 1500 and ws["G6"].value == "-" and ws["J6"].value == "릴스 없음"
    assert ws["E6"].hyperlink.target == "https://www.instagram.com/alpha/reels/"
    assert ws["F7"].value == "-" and ws["J7"].value == "실패 - 계정을 찾을 수 없습니다"
    assert ws["B9"].value.startswith("* 수집일시")

    feed = load_workbook(export_xlsx(tmp_path / "feed.xlsx", TITLES["feed"], COLUMNS["feed"], rows, notes("feed", 5))).active
    assert [feed.cell(row=4, column=c).value for c in range(2, 10)] == [
        "구분", "아이디", "이름", "프로필 링크", "팔로워", "평균 좋아요", "평균 댓글", "비고"]


def test_unknown_error_is_single_failure():
    client = fake_client({PROFILE + "x": Resp(500, "oops")})
    with pytest.raises(InstaError):
        client.profile("x")


# ---------------------------------------------------------------- 화면


def test_window_runs_and_exports(tmp_path, monkeypatch):
    import os
    import time

    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    monkeypatch.setenv("INSTA_CRAWLER_HOME", str(tmp_path))
    from PySide6.QtWidgets import QApplication, QFileDialog

    from insta_crawler import ui

    app = QApplication.instance() or QApplication([])
    from insta_crawler.theme import setup_app

    setup_app(app)
    client = fake_client({
        PROFILE + "alpha": Resp(200, profile("alpha", edges=[node("x", 2, 10, 2), node("y", 1, 30, 4, 900, True)])),
    })
    monkeypatch.setattr(ui, "InstaClient", lambda cookies: client)
    win = ui.MainWindow(Settings(request_delay=1, cooldown=0))
    win.settings.request_delay = 0

    def run_all():
        win.start_all()
        end = time.monotonic() + 10
        while (win.busy or win.thread is not None) and time.monotonic() < end:
            app.processEvents()
            time.sleep(0.02)

    win.edit.setPlainText("https://www.instagram.com/alpha/\n@nobody")
    assert win.count_chip.text() == "2개"
    assert win.mode == "feed" and win.mode_btns["feed"].isChecked()
    run_all()
    assert [r.status for r in win.model.rows] == ["완료", "실패"]
    assert win.model.rows[0].values["avg_likes"] == 20 and "avg_views" not in win.model.rows[0].values
    assert "완료 1" in win.summary.text() and win.retry_btn.isEnabled()
    assert win.model.tsv().splitlines()[0].split("\t") == ["No", "아이디", "이름", "팔로워", "평균 좋아요", "평균 댓글", "상태", "비고"]
    assert win.model.tsv().splitlines()[1].split("\t")[:3] == ["1", "@alpha", "앨리스"]

    win.mode_btns["reels"].click()  # 릴스로 바꾸면 표가 따로 (피드 결과는 그대로 남음)
    assert win.mode == "reels" and win.settings.mode == "reels" and win.model.rows == []
    assert "릴스" in win.desc.text()
    run_all()
    assert win.model.rows[0].values["avg_views"] == 900 and win.model.rows[0].values["avg_likes"] == 30
    assert win.model.tsv().splitlines()[0].split("\t")[3:7] == ["팔로워", "평균 조회수", "평균 좋아요", "평균 댓글"]
    assert len(win.models["feed"].rows) == 2

    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(tmp_path / "out.xlsx"), ""))
    monkeypatch.setattr(ui.QDesktopServices, "openUrl", lambda *a: True)
    win.export()
    ws = load_workbook(tmp_path / "out.xlsx").active
    assert ws["B2"].value == "인스타 릴스 지표 (2)" and ws["C6"].value == "alpha"
    assert ws["E6"].value == "https://www.instagram.com/alpha/reels/" and ws["E7"].value == "https://www.instagram.com/nobody/reels/"
    win.close()
