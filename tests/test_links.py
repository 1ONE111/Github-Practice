from openpyxl import Workbook

from naver_crawler.links import dedupe, load_links_file, read_text_file
from naver_crawler.parsing import extract_urls

MIXED = """1. 꼬마츄츄 http://blog.naver.com/baby0817 (요리) 연락: baby0817@naver.com
블로그: blog.naver.com/shriya, 카페 https://cafe.naver.com/joonggonara/123.
# 메모 줄은 무시 https://blog.naver.com/skip
단축 https://naver.me/abCD12 끝)  m.blog.naver.com/abc/22。 참고 ndex.kr/x
"""


def test_extract_urls_from_mixed_text():
    assert extract_urls(MIXED) == [
        "http://blog.naver.com/baby0817",
        "https://blog.naver.com/shriya",
        "https://cafe.naver.com/joonggonara/123",
        "https://naver.me/abCD12",
        "https://m.blog.naver.com/abc/22",
    ]


def test_extract_keeps_query_strings():
    url = "https://kin.naver.com/qna/detail.naver?d1id=8&dirId=80101&docId=123456"
    assert extract_urls(f"질문 {url} 입니다") == [url]


def test_txt_encodings(tmp_path):
    for i, enc in enumerate(("utf-8", "utf-8-sig", "cp949", "utf-16")):
        path = tmp_path / f"links{i}.txt"
        path.write_bytes(MIXED.encode(enc))
        assert "꼬마츄츄" in read_text_file(path)
        assert len(load_links_file(path)) == 5


def test_xlsx_reads_cells_and_hyperlinks(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["B6"] = "축복받은 블로그"
    ws["E6"] = "링크"
    ws["E6"].hyperlink = "http://blog.naver.com/baby0817"
    ws["E7"] = "https://blog.naver.com/shriya"
    ws["F7"] = "shriya@naver.com"
    wb.create_sheet("두번째")["A1"] = "kin.naver.com/qna/detail.naver?docId=1"
    path = tmp_path / "list.xlsx"
    wb.save(path)
    assert load_links_file(path) == [
        "http://blog.naver.com/baby0817",
        "https://blog.naver.com/shriya",
        "https://kin.naver.com/qna/detail.naver?docId=1",
    ]


def test_dedupe_against_existing():
    new, skipped = dedupe(["a", "b", "a", "c"], existing=["c"])
    assert new == ["a", "b"] and skipped == 2
