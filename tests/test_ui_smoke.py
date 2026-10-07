"""화면 없이(offscreen) 메인 창을 띄워 실제 수집 → 표 갱신 → 엑셀 저장까지 확인한다."""

import os
import time

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from tests.fake_naver import FakeHttp, FakeNaver  # noqa: E402

CHROME = os.environ.get("NC_TEST_CHROME")
DRIVER = os.environ.get("NC_TEST_DRIVER")
SHOTS = os.environ.get("NC_SCREENSHOT_DIR")


@pytest.fixture(scope="module")
def qapp():
    from PySide6.QtWidgets import QApplication

    from naver_crawler.ui.theme import STYLESHEET, app_font, apply_palette

    app = QApplication.instance() or QApplication([])
    app.setStyle("Fusion")
    apply_palette(app)
    app.setFont(app_font())
    app.setStyleSheet(STYLESHEET)
    return app


def wait_until(app, cond, timeout=60):
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        app.processEvents()
        if cond():
            return True
        time.sleep(0.05)
    return False


def test_window_builds(qapp, tmp_path, monkeypatch):
    monkeypatch.setenv("NAVER_CRAWLER_HOME", str(tmp_path))
    from naver_crawler.config import Settings
    from naver_crawler.ui.main_window import MainWindow

    win = MainWindow(Settings())
    from naver_crawler.ui.main_window import PAGE_ROLE

    selectable = [win.nav.item(i) for i in range(win.nav.count()) if win.nav.item(i).data(PAGE_ROLE) is not None]
    assert len(selectable) == 6
    win.show_page(4)
    assert win.current_page().task.id == "blog_post"
    assert win.page_title.text() == "블로그 포스팅"
    assert [p.task.id for p in win.pages] == ["cafe_info", "cafe_article", "cafe_avg", "blog_info", "blog_post", "kin"]
    win.close()


@pytest.mark.skipif(not (CHROME and DRIVER), reason="NC_TEST_CHROME / NC_TEST_DRIVER 미설정")
def test_crawl_and_export_through_ui(qapp, tmp_path, monkeypatch):
    monkeypatch.setenv("NAVER_CRAWLER_HOME", str(tmp_path / "home"))
    from naver_crawler.config import Settings
    from naver_crawler.ui.main_window import MainWindow

    with FakeNaver(tmp_path) as fake:
        monkeypatch.setenv("NAVER_CRAWLER_CHROME_ARGS", fake.chrome_args())
        settings = Settings(headless=True, keep_login=False, request_delay=0, page_timeout=5,
                            chromedriver_path=DRIVER, chrome_binary=CHROME)
        win = MainWindow(settings)
        win.resize(1480, 900)
        win.show()
        visitors = "".join(f'<visitorcnt id="2026100{i}" cnt="{c}"/>' for i, c in enumerate([10, 20, 30, 40, 50], 1))
        win.engine.http = FakeHttp({
            "https://blog.naver.com/NVisitorgp4Ajax.nhn?blogId=tester": f"<visitorcnts>{visitors}</visitorcnts>",
        })
        page = win.pages[3]
        win.show_page(3)
        links = tmp_path / "links.txt"
        links.write_bytes("메모\nblog.naver.com/tester 요리 블로그\n연락 tester@naver.com\n".encode("cp949"))
        page.load_files([str(links)])
        assert page.urls() == ["https://blog.naver.com/tester"]
        page.load_files([str(links)])  # 같은 파일 다시 → 중복은 건너뜀
        assert page.urls() == ["https://blog.naver.com/tester"]
        if SHOTS:
            os.makedirs(SHOTS, exist_ok=True)
            qapp.processEvents()
            win.grab().save(os.path.join(SHOTS, "main_empty.png"))
        page.url_edit.setPlainText("https://blog.naver.com/tester\nblog.naver.com/nowidget\n\nhttps://blog.naver.com/tester")
        assert page.url_count.text() == "3개"
        page.start_all()
        assert win.busy and not page.start_btn.isEnabled()
        assert wait_until(qapp, lambda: not win.busy and win.thread is None, timeout=90)

        statuses = [r.status for r in page.model.rows]
        assert statuses == ["완료", "위젯없음", "완료"]
        assert page.model.rows[0].values["avg"] == 30
        assert "위젯없음 1" in page.summary.text()
        assert page.start_btn.isEnabled()

        # 카페 게시글 탭도 하나 채워 전체 저장 확인
        cafe = win.pages[1]
        cafe.url_edit.setPlainText("https://cafe.naver.com/testcafe/123")
        cafe.start_all()
        assert wait_until(qapp, lambda: not win.busy and win.thread is None, timeout=60)
        assert cafe.model.rows[0].status == "완료"

        if SHOTS:
            os.makedirs(SHOTS, exist_ok=True)
            win.show_page(3)
            win.log_toggle.setChecked(True)
            qapp.processEvents()
            win.grab().save(os.path.join(SHOTS, "main_blog.png"))
            win.log_toggle.setChecked(False)
            win.show_page(1)
            qapp.processEvents()
            win.grab().save(os.path.join(SHOTS, "main_cafe.png"))

        from naver_crawler.excel import export_xlsx
        from openpyxl import load_workbook

        out = export_xlsx(tmp_path / "all.xlsx", [p.model.sheet_data() for p in win.pages if p.model.rows])
        if SHOTS:
            import shutil

            shutil.copy(out, os.path.join(SHOTS, "sample.xlsx"))
        wb = load_workbook(out)
        assert wb.sheetnames == ["카페 게시글", "블로그 정보", "제외 목록"]
        win.close()
