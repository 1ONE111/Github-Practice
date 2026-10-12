"""링크 목록 파일 읽기(txt / csv / xlsx) + 결과 엑셀 저장 (팀 리스트 양식)."""

from __future__ import annotations

import datetime as dt
import re
import unicodedata
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from insta_crawler.collector import STATUS_CANCELLED, STATUS_FAIL, STATUS_OK, STATUS_PARTIAL, STATUS_PRIVATE, Column

SUPPORTED = (".txt", ".csv", ".tsv", ".xlsx", ".xlsm")


def read_links_text(path: str | Path) -> str:
    """파일 내용을 텍스트로 (엑셀은 셀 값 · 하이퍼링크를 한 줄씩). 링크 추출은 parse_targets 가 한다."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        wb = load_workbook(path, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.hyperlink is not None and cell.hyperlink.target:
                        lines.append(str(cell.hyperlink.target))
                    elif cell.value is not None:
                        lines.append(str(cell.value))
        return "\n".join(lines)
    data = path.read_bytes()
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    for encoding in ("utf-8-sig", "cp949"):  # 메모장 ANSI 저장 파일 대응
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


# ---------------------------------------------------------------- 엑셀 저장
# B열 시작, B2 검정 제목줄 '제목 (건수)', 4행 검정 머리글, 6행부터 데이터, 8pt · 얇은 테두리, 표 아래 '*' 주석

HEAD_FONT = "나눔고딕"
BODY_FONT = "맑은 고딕"
SIZE = 8
NUMBER_FORMAT = "#,##0\\ "
_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BLACK = PatternFill("solid", fgColor="000000")
_LEFT = Alignment(horizontal="left", vertical="center")


def note_for(row: dict) -> str:
    status, message = row.get("status"), (row.get("message") or "").strip()
    if status in (STATUS_OK, STATUS_PARTIAL, STATUS_PRIVATE):
        return message or ("" if status == STATUS_OK else str(status))
    if status == STATUS_FAIL:
        return f"실패 - {message}" if message else "실패"
    if status == STATUS_CANCELLED:
        return "중지됨 (미수집)"
    return "미수집"


def _width(text: str) -> float:
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in text)


def export_xlsx(path: str | Path, title: str, columns: Sequence[Column], rows: list[dict],
                notes: Sequence[str] = (), generated: dt.datetime | None = None) -> Path:
    """rows: {"url", "status", "message", <열 key>…}. 실패한 행도 '-' 와 비고로 남긴다."""
    generated = generated or dt.datetime.now()
    cols = [Column("_no", "구분", "int", 6), *columns, Column("_note", "비고", width=30)]
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", " ", title)[:31] or "결과"
    ws.column_dimensions["A"].width = 2.8
    first, last = 2, 1 + len(cols)

    ws.merge_cells(start_row=2, start_column=first, end_row=2, end_column=last)
    cell = ws.cell(row=2, column=first, value=f"{title} ({len(rows)})")
    cell.font = Font(name=HEAD_FONT, size=SIZE, bold=True, color="FFFFFF")
    cell.fill = _BLACK
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, col in enumerate(cols):
        cell = ws.cell(row=4, column=first + i, value=col.header)
        cell.font = Font(name=HEAD_FONT, size=SIZE, color="FFFFFF")
        cell.fill = _BLACK
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER
    ws.row_dimensions[5].height = 3.75

    widths = [max(col.width, _width(col.header) + 2) for col in cols]
    for r, row in enumerate(rows):
        excel_row = 6 + r
        ws.row_dimensions[excel_row].height = 18
        for i, col in enumerate(cols):
            cell = ws.cell(row=excel_row, column=first + i)
            cell.font = Font(name=BODY_FONT, size=SIZE)
            cell.border = _BORDER
            cell.alignment = _LEFT
            value = r + 1 if col.key == "_no" else note_for(row) if col.key == "_note" else row.get(col.key)
            if col.kind == "int":
                if isinstance(value, int) and not isinstance(value, bool):
                    cell.value = value
                    cell.number_format = NUMBER_FORMAT
                else:
                    cell.value = "-"
            else:
                cell.value = "-" if value in (None, "") and col.key != "_note" else (value or "")
                cell.number_format = "@"
                if col.key == "url" and value:
                    cell.hyperlink = value
                    cell.font = Font(name=BODY_FONT, size=SIZE, color="0000FF", underline="single")
            text = f"{cell.value:,} " if isinstance(cell.value, int) else str(cell.value)
            widths[i] = max(widths[i], min(_width(text) * 0.9 + 2, 45))
    for i, width in enumerate(widths):
        ws.column_dimensions[get_column_letter(first + i)].width = round(width, 1)

    for k, note in enumerate([f"* 수집일시: {generated:%Y-%m-%d %H:%M}", *notes]):
        cell = ws.cell(row=6 + len(rows) + 1 + k, column=first, value=note)
        cell.font = Font(name=BODY_FONT, size=SIZE)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.properties.creator = "인스타 크롤러"

    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    wb.save(path)
    return path


def default_filename(label: str, when: dt.datetime | None = None) -> str:
    return f"인스타크롤링_{label}_{(when or dt.datetime.now()):%y%m%d}.xlsx"
